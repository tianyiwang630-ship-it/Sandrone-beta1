from __future__ import annotations

import os
import mimetypes
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse

from agent.server.models import (
    CreateFolderRequest,
    FileConflictCheckRequest,
    FileConflictCheckResponse,
    FileContentResponse,
    FileInfo,
    FileListResponse,
    OpenInFolderRequest,
    UploadConflictItem,
)
from agent.server.deps import state_store

router = APIRouter(prefix="/api/files", tags=["files"])

TEXT_EXTENSIONS = {
    ".bat",
    ".css",
    ".csv",
    ".html",
    ".js",
    ".json",
    ".jsx",
    ".log",
    ".md",
    ".ps1",
    ".py",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".yaml",
    ".yml",
}
PREVIEW_LIMIT_BYTES = 256 * 1024
DOCUMENT_PREVIEW_CHAR_LIMIT = 50_000
UPLOAD_LIMIT_BYTES = 100 * 1024 * 1024


def _project_root(project_id: str) -> Path:
    project = state_store.get_project(project_id)
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")
    return Path(project["workspace_path"]).resolve()


def _resolve_child(root: Path, relative_path: str | None) -> Path:
    rel = (relative_path or "").strip().replace("\\", "/")
    target = (root / rel).resolve()
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Path is outside the project workspace") from exc
    return target


def _resolve_target_dir(root: Path, relative_path: str | None) -> Path:
    target = _resolve_child(root, relative_path)
    if not target.exists():
        raise HTTPException(status_code=404, detail="Directory not found")
    if not target.is_dir():
        raise HTTPException(status_code=400, detail="Target path is not a directory")
    return target


def _to_relative(root: Path, path: Path) -> str:
    return path.resolve().relative_to(root.resolve()).as_posix()


def _sanitize_folder_name(name: str) -> str:
    cleaned = name.strip().replace("\\", "_").replace("/", "_")
    if not cleaned or cleaned in {".", ".."}:
        raise HTTPException(status_code=400, detail="Invalid folder name")
    invalid = '<>:"|?*'
    for char in invalid:
        cleaned = cleaned.replace(char, "_")
    return cleaned.rstrip(". ").strip() or "new-folder"


def _split_name(name: str) -> tuple[str, str]:
    suffix = "".join(Path(name).suffixes)
    if not suffix:
        return name, ""
    return name[: -len(suffix)], suffix


def _rename_candidate(path: Path) -> Path:
    base, ext = _split_name(path.name)
    copy_match = re.fullmatch(r"(.*)（副本(?:(\d+))?）", base)
    if copy_match:
        base = copy_match.group(1)
        counter = int(copy_match.group(2) or 1) + 1
    else:
        counter = 1
    while True:
        copy_suffix = "（副本）" if counter == 1 else f"（副本{counter}）"
        candidate = path.with_name(f"{base}{copy_suffix}{ext}")
        if not candidate.exists():
            return candidate
        counter += 1


def _apply_conflict_strategy(dest: Path, strategy: str | None) -> Path:
    if not dest.exists():
        return dest
    if strategy == "replace":
        if dest.is_dir():
            shutil.rmtree(dest)
        else:
            dest.unlink()
        return dest
    if strategy == "rename":
        return _rename_candidate(dest)
    raise HTTPException(status_code=409, detail=f"Target already exists: {dest.name}")


def _collect_conflicts(root: Path, target_path: str, relative_paths: list[str]) -> list[UploadConflictItem]:
    target_dir = _resolve_target_dir(root, target_path)
    conflicts: list[UploadConflictItem] = []
    seen: set[str] = set()
    for relative_path in relative_paths:
        final_path = _resolve_child(target_dir, relative_path)
        if not final_path.exists():
            continue
        relative = _to_relative(root, final_path)
        if relative in seen:
            continue
        seen.add(relative)
        conflicts.append(
            UploadConflictItem(
                path=relative,
                name=final_path.name,
                is_dir=final_path.is_dir(),
            )
        )
    return conflicts


def _list_dir(root: Path, directory: Path) -> list[FileInfo]:
    result: list[FileInfo] = []
    for entry in sorted(directory.iterdir(), key=lambda item: (not item.is_dir(), item.name.lower())):
        if entry.name in {".git", "__pycache__", "node_modules"}:
            continue
        stat = entry.stat()
        result.append(
            FileInfo(
                name=entry.name,
                path=_to_relative(root, entry),
                size=stat.st_size,
                is_dir=entry.is_dir(),
            )
        )
    return result


def _write_upload_file(upload: UploadFile, dest: Path) -> int:
    total = 0
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            delete=False,
            dir=str(dest.parent),
            prefix=f".{dest.name}.",
            suffix=".uploading",
        ) as handle:
            temp_path = Path(handle.name)
            while True:
                chunk = upload.file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > UPLOAD_LIMIT_BYTES:
                    raise HTTPException(status_code=413, detail="File too large, max 100MB")
                handle.write(chunk)
        if temp_path is None:
            raise RuntimeError("Temporary upload file was not created")
        temp_path.replace(dest)
        return total
    finally:
        if temp_path and temp_path.exists():
            try:
                temp_path.unlink()
            except Exception:
                pass
        upload.file.close()


def _truncate_document_preview(content: str) -> str:
    if len(content) <= DOCUMENT_PREVIEW_CHAR_LIMIT:
        return content
    return f"{content[:DOCUMENT_PREVIEW_CHAR_LIMIT].rstrip()}\n\n[预览已截断，文件内容更长。]"


def _file_response(
    root: Path,
    target: Path,
    size: int,
    language: str,
    content: str | None,
    previewable: bool,
    message: str | None = None,
) -> FileContentResponse:
    return FileContentResponse(
        path=_to_relative(root, target),
        name=target.name,
        size=size,
        language=language,
        content=content,
        previewable=previewable,
        message=message,
    )


def _docx_table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()).replace("|", "\\|") for cell in row.cells]
        if any(cells):
            rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _preview_docx(root: Path, target: Path, size: int) -> FileContentResponse:
    try:
        from docx import Document
    except ImportError:
        return _file_response(root, target, size, "docx", None, False, "当前环境缺少 python-docx，暂时不能预览 Word 文档。")

    try:
        document = Document(str(target))
    except Exception as exc:
        return _file_response(root, target, size, "docx", None, False, f"Word 文档读取失败：{exc}")

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for index, table in enumerate(document.tables, start=1):
        table_text = _docx_table_to_text(table)
        if table_text:
            blocks.append(f"表格 {index}\n{table_text}")

    content = "\n\n".join(blocks).strip()
    if not content:
        return _file_response(root, target, size, "docx", None, False, "Word 文档没有提取到可预览文本。")
    return _file_response(root, target, size, "markdown", _truncate_document_preview(content), True)


def _preview_pdf(root: Path, target: Path, size: int) -> FileContentResponse:
    return _file_response(root, target, size, "pdf", None, True, "PDF 将使用浏览器内置查看器预览。")


def _preview_pptx(root: Path, target: Path, size: int) -> FileContentResponse:
    return _file_response(root, target, size, "pptx", None, True, "PPTX 将在文件预览区显示幻灯片。")


def _asset_media_type(target: Path) -> str | None:
    media_type, _ = mimetypes.guess_type(target.name)
    return media_type


@router.get("/tree", response_model=FileListResponse)
def list_files(project_id: str, path: str | None = Query(default=None)):
    root = _project_root(project_id)
    target = _resolve_target_dir(root, path or "")
    return FileListResponse(root=str(root), path=_to_relative(root, target) if target != root else "", items=_list_dir(root, target))


@router.get("/content", response_model=FileContentResponse)
def read_file(project_id: str, path: str):
    root = _project_root(project_id)
    target = _resolve_child(root, path)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    size = target.stat().st_size
    suffix = target.suffix.lower()
    if suffix == ".docx":
        return _preview_docx(root, target, size)
    if suffix == ".pdf":
        return _preview_pdf(root, target, size)
    if suffix == ".pptx":
        return _preview_pptx(root, target, size)
    if size > PREVIEW_LIMIT_BYTES:
        return _file_response(root, target, size, target.suffix.lstrip(".") or "text", None, False, "文件太大，第一版只预览 256KB 以内的文本文件。")
    if suffix not in TEXT_EXTENSIONS:
        return _file_response(root, target, size, target.suffix.lstrip(".") or "file", None, False, "这个文件类型第一版暂不支持预览。")
    try:
        content = target.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = target.read_text(encoding="gb18030", errors="replace")
    return _file_response(root, target, size, target.suffix.lstrip(".") or "text", content, True)


@router.get("/raw")
def raw_file(project_id: str, path: str):
    root = _project_root(project_id)
    target = _resolve_child(root, path)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    suffix = target.suffix.lower()
    media_types = {
        ".pdf": "application/pdf",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    if suffix not in media_types:
        raise HTTPException(status_code=400, detail="Only PDF and PPTX inline preview are supported")
    return FileResponse(str(target), media_type=media_types[suffix])


@router.get("/assets/{project_id}/{asset_path:path}")
def file_asset(project_id: str, asset_path: str):
    root = _project_root(project_id)
    target = _resolve_child(root, asset_path)
    if not target.exists() or not target.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(str(target), media_type=_asset_media_type(target))


@router.post("/conflicts/check", response_model=FileConflictCheckResponse)
def check_conflicts(body: FileConflictCheckRequest):
    root = _project_root(body.project_id)
    conflicts = _collect_conflicts(root, body.target_path, body.relative_paths)
    return FileConflictCheckResponse(has_conflicts=bool(conflicts), conflicts=conflicts)


@router.post("/upload")
def upload_file(
    file: UploadFile = File(...),
    project_id: str = Form(...),
    target_path: str = Form(default=""),
    relative_path: str = Form(default=""),
    conflict_strategy: str | None = Form(default=None),
):
    root = _project_root(project_id)
    target_dir = _resolve_target_dir(root, target_path)
    relative = (relative_path or file.filename or "").strip().replace("\\", "/")
    if not relative:
        raise HTTPException(status_code=400, detail="Missing target filename")
    dest = _resolve_child(target_dir, relative)
    dest.parent.mkdir(parents=True, exist_ok=True)
    final_dest = _apply_conflict_strategy(dest, conflict_strategy)
    final_dest.parent.mkdir(parents=True, exist_ok=True)
    size = _write_upload_file(file, final_dest)
    return FileInfo(name=final_dest.name, path=_to_relative(root, final_dest), size=size, is_dir=False)


@router.post("/folders")
def create_folder(body: CreateFolderRequest):
    root = _project_root(body.project_id)
    target_dir = _resolve_target_dir(root, body.target_path)
    folder_name = _sanitize_folder_name(body.name)
    folder_path = _resolve_child(target_dir, folder_name)
    if folder_path.exists():
        raise HTTPException(status_code=409, detail="Folder already exists")
    folder_path.mkdir(parents=False, exist_ok=False)
    return FileInfo(name=folder_path.name, path=_to_relative(root, folder_path), size=0, is_dir=True)


@router.post("/open-in-folder")
def open_in_folder(body: OpenInFolderRequest):
    root = _project_root(body.project_id)
    if not body.path:
        os.startfile(str(root))
        return {"ok": True}

    target = _resolve_child(root, body.path)
    if not target.exists():
        raise HTTPException(status_code=404, detail="Path not found")

    subprocess.Popen(["explorer.exe", "/select,", str(target)])
    return {"ok": True}
