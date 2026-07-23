#!/usr/bin/env python3
"""
意识启发的 Agent 架构 → PPTX
简短、紧凑、视觉清晰
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Colors ──
INK       = RGBColor(0x1a, 0x2a, 0x4a)  # deep navy
BLUE      = RGBColor(0x24, 0x57, 0xd6)
BLUE_LIGHT = RGBColor(0xdf, 0xe8, 0xff)
CORAL     = RGBColor(0xdf, 0x65, 0x4c)
CORAL_LIGHT = RGBColor(0xf8, 0xdf, 0xd9)
GREEN     = RGBColor(0x18, 0x79, 0x6b)
GREEN_LIGHT = RGBColor(0xd9, 0xee, 0xe9)
WHITE     = RGBColor(0xff, 0xff, 0xff)
NEAR_WHITE = RGBColor(0xf7, 0xf4, 0xed)
MUTED     = RGBColor(0x64, 0x70, 0x80)
DARK_BG   = RGBColor(0x19, 0x22, 0x2e)
DARK_TEXT  = RGBColor(0xf5, 0xf1, 0xe8)
DARK_MUTED = RGBColor(0xae, 0xb7, 0xc3)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height

# ── Helper functions ──

def add_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    return shape

def add_rounded_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    return shape

def set_text(shape, text, size=14, color=INK, bold=False, alignment=PP_ALIGN.LEFT, font_name='Arial'):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font_name
    return p

def add_text_box(slide, left, top, width, height, text, size=14, color=INK, bold=False, alignment=PP_ALIGN.LEFT, font_name='Arial'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font_name
    return txBox

def add_multi_text(slide, left, top, width, height, lines, font_name='Arial'):
    """lines: list of (text, size, color, bold, alignment)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold, alignment) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = alignment
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font_name
    return txBox

def add_card(slide, left, top, w, h, title, body, title_color=INK, body_color=MUTED, bg=WHITE, border=None):
    card = add_rounded_rect(slide, left, top, w, h, fill_color=bg, border_color=border or RGBColor(0xe0, 0xe0, 0xe0))
    # Title
    add_text_box(slide, left + Inches(0.2), top + Inches(0.15), w - Inches(0.4), Inches(0.5),
                 title, size=13, color=title_color, bold=True)
    # Body
    add_text_box(slide, left + Inches(0.2), top + Inches(0.6), w - Inches(0.4), h - Inches(0.75),
                 body, size=10, color=body_color)
    return card

# ══════════════════════════════════════════════════════════════
# SLIDE 1 — 标题页
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# 背景色块 - 左半深色
add_rect(slide, Inches(0), Inches(0), Inches(6.8), H, fill_color=INK)
# 右侧浅色
add_rect(slide, Inches(6.8), Inches(0), Inches(6.6), H, fill_color=NEAR_WHITE)

# 左侧内容
add_text_box(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(1.2),
             "CONSCIOUSNESS-INSPIRED\nAGENT ARCHITECTURE",
             size=32, color=WHITE, bold=True)

# 装饰线
add_rect(slide, Inches(0.8), Inches(3.1), Inches(2.0), Pt(3), fill_color=BLUE)

add_text_box(slide, Inches(0.8), Inches(3.4), Inches(5.5), Inches(1.5),
             "让 Agent 形成一个持续的自己\n—— 从意识科学中提取工程机制",
             size=18, color=DARK_MUTED)

add_text_box(slide, Inches(0.8), Inches(5.5), Inches(5.5), Inches(0.5),
             "概念架构 · 未来研究方向", size=11, color=DARK_MUTED)

# 右侧 - 核心词云
keywords = ["HOT 二阶表征", "元认知", "自传体记忆", "DMN", "双脑机制", "Observer", "做梦机制", "主体连续性"]
positions = [
    (Inches(7.5), Inches(1.8)),
    (Inches(10.5), Inches(1.5)),
    (Inches(8.0), Inches(3.2)),
    (Inches(11.0), Inches(3.0)),
    (Inches(7.5), Inches(4.5)),
    (Inches(10.2), Inches(4.8)),
    (Inches(8.5), Inches(5.8)),
    (Inches(11.2), Inches(5.5)),
]
for kw, (x, y) in zip(keywords, positions):
    add_text_box(slide, x, y, Inches(2.5), Inches(0.4), kw, size=11, color=BLUE, bold=True)


# ══════════════════════════════════════════════════════════════
# SLIDE 2 — 核心主张
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, Inches(0), Inches(0), W, H, fill_color=NEAR_WHITE)

# 左侧编号
add_text_box(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(0.4),
             "01", size=14, color=BLUE, bold=True)

# 核心主张标题
add_text_box(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.8),
             "核 心 主 张", size=28, color=INK, bold=True)
add_rect(slide, Inches(0.6), Inches(1.8), Inches(1.5), Pt(3), fill_color=CORAL)

# 主张正文
add_text_box(slide, Inches(0.6), Inches(2.3), Inches(12), Inches(1.5),
             "把 Agent 从一次性任务执行器，推进为能够观察自身、理解场景，\n并在长期交互中形成连续性的协作主体。",
             size=18, color=INK)

# 两个对比框
# 左边：传统
box1 = add_rounded_rect(slide, Inches(0.6), Inches(4.2), Inches(5.5), Inches(2.5),
                        fill_color=WHITE, border_color=RGBColor(0xdd, 0xdd, 0xdd))
add_text_box(slide, Inches(0.9), Inches(4.4), Inches(5.0), Inches(0.4),
             "❌  传统 Agent", size=16, color=MUTED, bold=True)
add_text_box(slide, Inches(0.9), Inches(4.9), Inches(5.0), Inches(1.5),
             "• 一次性任务执行器\n• 每次交互从零开始\n• 无自身认知和连续性\n• 遇到失败易重复相同路径",
             size=12, color=MUTED)

# 右边：目标
box2 = add_rounded_rect(slide, Inches(6.5), Inches(4.2), Inches(6.2), Inches(2.5),
                        fill_color=BLUE_LIGHT, border_color=BLUE)
add_text_box(slide, Inches(6.8), Inches(4.4), Inches(5.5), Inches(0.4),
             "✅  意识启发 Agent", size=16, color=BLUE, bold=True)
add_text_box(slide, Inches(6.8), Inches(4.9), Inches(5.5), Inches(1.5),
             "• 拥有对自身状态的二阶观察\n• 从历史中形成持续的自我理解\n• 场景感知，个性化响应\n• 从失败中提取元认知经验",
             size=12, color=INK)

# 底部箭条
add_text_box(slide, Inches(4.5), Inches(6.9), Inches(4.5), Inches(0.4),
             "目标：主体连续性 → 长期协作主体",
             size=12, color=BLUE, bold=True, alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# SLIDE 3 — 理论借鉴：四个可转译机制
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, Inches(0), Inches(0), W, H, fill_color=NEAR_WHITE)

add_text_box(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(0.4),
             "02", size=14, color=BLUE, bold=True)
add_text_box(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.6),
             "从意识理论中提取工程机制", size=26, color=INK, bold=True)
add_rect(slide, Inches(0.6), Inches(1.55), Inches(1.5), Pt(3), fill_color=CORAL)
add_text_box(slide, Inches(0.6), Inches(1.8), Inches(12), Inches(0.5),
             "不是复刻意识，而是借用四个可转译的认知机制，为长期 Agent 提供新架构思路",
             size=12, color=MUTED)

# 四张卡片
cards_data = [
    ("01", "HOT 二阶表征", "系统不仅处理任务，还能把\n自己的处理状态作为观察对象。\n\n→ Agent 拥有对自身思考\n   过程的元表示"),
    ("02", "元认知", "评估当前判断、方法\n与能力边界。\n\n→ 决定何时坚持、修正\n   或求助"),
    ("03", "自传体记忆", "经历不只是日志，而是构成\n\"我过去怎样行动\"的\n连续历史。\n\n→ 形成主体连续性"),
    ("04", "DMN\n默认模式网络", "在离线状态重组经历、\n自我与关系。\n\n→ 为下一次行动形成\n   背景理解"),
]

card_w = Inches(2.85)
card_h = Inches(4.8)
gap = Inches(0.25)
start_x = Inches(0.6)
start_y = Inches(2.5)

for i, (num, title, body) in enumerate(cards_data):
    x = start_x + i * (card_w + gap)
    y = start_y
    
    # Card background
    card = add_rounded_rect(slide, x, y, card_w, card_h, fill_color=WHITE,
                            border_color=RGBColor(0xe5, 0xe5, 0xe5))
    
    # Number
    add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(0.6), Inches(0.4),
                 num, size=14, color=BLUE, bold=True)
    
    # Title
    add_text_box(slide, x + Inches(0.2), y + Inches(0.8), card_w - Inches(0.4), Inches(0.8),
                 title, size=15, color=INK, bold=True)
    
    # Body
    add_text_box(slide, x + Inches(0.2), y + Inches(1.8), card_w - Inches(0.4), card_h - Inches(2.2),
                 body, size=10, color=MUTED)


# ══════════════════════════════════════════════════════════════
# SLIDE 4 — 双脑机制
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, Inches(0), Inches(0), W, H, fill_color=NEAR_WHITE)

add_text_box(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(0.4),
             "03", size=14, color=BLUE, bold=True)
add_text_box(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.6),
             "双脑机制：行动脑 + 主体脑", size=26, color=INK, bold=True)
add_rect(slide, Inches(0.6), Inches(1.55), Inches(1.5), Pt(3), fill_color=CORAL)

# 关键陈述
add_text_box(slide, Inches(0.6), Inches(1.9), Inches(12), Inches(0.8),
             "普通 Agent 只问\"现在怎么做\"。双脑机制还会追问：结合我是谁、用户是谁、过去发生过什么，这件事应该以什么方式做？",
             size=12, color=CORAL, bold=True)

# 三列架构图
col_w = Inches(3.6)
col_gap = Inches(0.35)
col_start_y = Inches(2.8)

# 左列 - 输入
left_x = Inches(0.6)
add_text_box(slide, left_x, col_start_y - Inches(0.35), col_w, Inches(0.3),
             "外部输入", size=9, color=MUTED, bold=True)

input_box = add_rounded_rect(slide, left_x, col_start_y, col_w, Inches(1.2),
                             fill_color=WHITE, border_color=RGBColor(0xdd, 0xdd, 0xdd))
add_text_box(slide, left_x + Inches(0.15), col_start_y + Inches(0.1), col_w - Inches(0.3), Inches(0.4),
             "用户与当前任务", size=13, color=INK, bold=True)
add_text_box(slide, left_x + Inches(0.15), col_start_y + Inches(0.55), col_w - Inches(0.3), Inches(0.5),
             "需求、文件、环境、即时反馈", size=10, color=MUTED)

# 箭头
add_text_box(slide, left_x, col_start_y + Inches(1.3), col_w, Inches(0.3),
             "↓", size=18, color=MUTED, alignment=PP_ALIGN.CENTER)

mem_box = add_rounded_rect(slide, left_x, col_start_y + Inches(1.7), col_w, Inches(1.1),
                           fill_color=GREEN_LIGHT, border_color=RGBColor(0xaa, 0xdd, 0xd5))
add_text_box(slide, left_x + Inches(0.15), col_start_y + Inches(1.8), col_w - Inches(0.3), Inches(0.35),
             "历史背景", size=13, color=GREEN, bold=True)
add_text_box(slide, left_x + Inches(0.15), col_start_y + Inches(2.2), col_w - Inches(0.3), Inches(0.5),
             "用户偏好 · 共同经历 · 项目约束", size=10, color=GREEN)

# 中列 - 行动脑
mid_x = Inches(4.9)
add_text_box(slide, mid_x, col_start_y - Inches(0.35), col_w, Inches(0.3),
             "一阶 · 行动脑", size=9, color=MUTED, bold=True)

wa_box = add_rounded_rect(slide, mid_x, col_start_y, col_w, Inches(1.2),
                          fill_color=BLUE, border_color=BLUE)
add_text_box(slide, mid_x + Inches(0.15), col_start_y + Inches(0.1), col_w - Inches(0.3), Inches(0.4),
             "Working Agent", size=13, color=WHITE, bold=True)
add_text_box(slide, mid_x + Inches(0.15), col_start_y + Inches(0.55), col_w - Inches(0.3), Inches(0.5),
             "理解任务、调用工具、做出选择并交付结果", size=10, color=RGBColor(0xcc, 0xdd, 0xff))

# 双向箭头
add_text_box(slide, mid_x, col_start_y + Inches(1.2), col_w, Inches(0.4),
             "↕", size=20, color=MUTED, alignment=PP_ALIGN.CENTER)

trail_box = add_rounded_rect(slide, mid_x, col_start_y + Inches(1.7), col_w, Inches(1.1),
                             fill_color=WHITE, border_color=RGBColor(0xdd, 0xdd, 0xdd))
add_text_box(slide, mid_x + Inches(0.15), col_start_y + Inches(1.8), col_w - Inches(0.3), Inches(0.35),
             "任务轨迹", size=13, color=INK, bold=True)
add_text_box(slide, mid_x + Inches(0.15), col_start_y + Inches(2.2), col_w - Inches(0.3), Inches(0.5),
             "目标 · 关键决策 · 异常 · 结果", size=10, color=MUTED)

# 右列 - 主体脑
right_x = Inches(9.2)
add_text_box(slide, right_x, col_start_y - Inches(0.35), col_w, Inches(0.3),
             "二阶 · 主体脑", size=9, color=MUTED, bold=True)

obs_box = add_rounded_rect(slide, right_x, col_start_y, col_w, Inches(1.2),
                           fill_color=CORAL_LIGHT, border_color=RGBColor(0xf0, 0xcc, 0xc5))
add_text_box(slide, right_x + Inches(0.15), col_start_y + Inches(0.1), col_w - Inches(0.3), Inches(0.4),
             "Observer Agent", size=13, color=CORAL, bold=True)
add_text_box(slide, right_x + Inches(0.15), col_start_y + Inches(0.55), col_w - Inches(0.3), Inches(0.5),
             "以行动脑为观察对象，检索历史，输出认知提示", size=10, color=CORAL)

# 箭头
add_text_box(slide, right_x, col_start_y + Inches(1.2), col_w, Inches(0.3),
             "↓", size=18, color=MUTED, alignment=PP_ALIGN.CENTER)

meta_box = add_rounded_rect(slide, right_x, col_start_y + Inches(1.7), col_w, Inches(1.1),
                            fill_color=WHITE, border_color=RGBColor(0xdd, 0xdd, 0xdd))
add_text_box(slide, right_x + Inches(0.15), col_start_y + Inches(1.8), col_w - Inches(0.3), Inches(0.35),
             "元认知提示", size=13, color=INK, bold=True)
add_text_box(slide, right_x + Inches(0.15), col_start_y + Inches(2.2), col_w - Inches(0.3), Inches(0.5),
             "场景框架 · 偏好提醒 · 风险与纠偏", size=10, color=MUTED)

# 底部的四个职责标签
jobs = [("场景化", "识别任务类型"), ("个性化", "召回用户偏好"),
        ("元认知", "观察偏离与失败"), ("连续性", "接续过去的自己")]
job_x = [Inches(0.6), Inches(3.7), Inches(6.8), Inches(9.9)]
for i, (title, desc) in enumerate(jobs):
    x = job_x[i]
    job_bg = add_rounded_rect(slide, x, Inches(5.5), Inches(2.8), Inches(0.8),
                              fill_color=CORAL_LIGHT, border_color=RGBColor(0xf0, 0xcc, 0xc5))
    add_text_box(slide, x + Inches(0.1), Inches(5.55), Inches(1.8), Inches(0.35),
                 title, size=11, color=CORAL, bold=True)
    add_text_box(slide, x + Inches(1.5), Inches(5.55), Inches(1.2), Inches(0.35),
                 desc, size=9, color=MUTED)

# 底部场景提示
add_text_box(slide, Inches(0.6), Inches(6.5), Inches(12), Inches(0.6),
             "场景示例：写 PRD → 识别设计场景，检索偏好，先讨论边界再落笔  |  连续失败 → 召回类似失败，判断问题根源，换方向尝试",
             size=10, color=MUTED)


# ══════════════════════════════════════════════════════════════
# SLIDE 5 — 做梦机制：离线重组
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
# 深色背景
add_rect(slide, Inches(0), Inches(0), W, H, fill_color=DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(0.4),
             "04", size=14, color=RGBColor(0x87, 0xa6, 0xff), bold=True)
add_text_box(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.6),
             "做梦机制：离线重组经历", size=26, color=DARK_TEXT, bold=True)
add_rect(slide, Inches(0.6), Inches(1.55), Inches(1.5), Pt(3), fill_color=RGBColor(0x70, 0xc7, 0xb8))

add_text_box(slide, Inches(0.6), Inches(1.9), Inches(12), Inches(0.8),
             "普通总结回答\"以后该遵守什么规则\"。做梦机制进一步回答：这些经历如何改变了我对自己、用户和双方关系的理解？",
             size=12, color=RGBColor(0x70, 0xc7, 0xb8), bold=True)

# 三列流程
# 左列 - 经历输入
add_text_box(slide, Inches(0.6), Inches(2.8), Inches(3.5), Inches(0.3),
             "经历输入", size=9, color=DARK_MUTED, bold=True)

logs = [
    ("USER · 提出一个模糊方向", ""),
    ("AGENT · 直接开始输出", ""),
    ("FEEDBACK · 希望先讨论边界", ""),
    ("RESULT · 对话后方案更准确", ""),
]
for i, (log, _) in enumerate(logs):
    y = Inches(3.2) + i * Inches(0.55)
    log_bg = add_rounded_rect(slide, Inches(0.6), y, Inches(3.5), Inches(0.45),
                              fill_color=RGBColor(0x22, 0x2e, 0x3c),
                              border_color=RGBColor(0x33, 0x3f, 0x50))
    add_text_box(slide, Inches(0.75), y + Inches(0.05), Inches(3.2), Inches(0.35),
                 log, size=9, color=DARK_MUTED)

# 中列 - 离线重组
add_text_box(slide, Inches(4.7), Inches(2.8), Inches(4.0), Inches(0.3),
             "离线重组", size=9, color=DARK_MUTED, bold=True)

steps = [
    ("① 回放", "找出关键选择、反馈和情绪信号"),
    ("② 关联", "连接相似经历，区分偶然与稳定模式"),
    ("③ 主体化", "形成关于用户、关系与自身角色的解释"),
    ("④ 沉淀", "保留来源与可信度，写入长期记忆"),
]
for i, (title, desc) in enumerate(steps):
    y = Inches(3.2) + i * Inches(0.9)
    step_bg = add_rounded_rect(slide, Inches(4.7), y, Inches(4.0), Inches(0.8),
                               fill_color=RGBColor(0x22, 0x2e, 0x3c))
    # Blue left border via thin rect
    add_rect(slide, Inches(4.7), y, Pt(3), Inches(0.8), fill_color=RGBColor(0x87, 0xa6, 0xff))
    add_text_box(slide, Inches(4.9), y + Inches(0.05), Inches(3.6), Inches(0.3),
                 title, size=11, color=DARK_TEXT, bold=True)
    add_text_box(slide, Inches(4.9), y + Inches(0.35), Inches(3.6), Inches(0.35),
                 desc, size=9, color=DARK_MUTED)

# 右列 - 记忆产出
add_text_box(slide, Inches(9.3), Inches(2.8), Inches(3.5), Inches(0.3),
             "记忆产出", size=9, color=DARK_MUTED, bold=True)

mems = [
    ("情节记忆", "某次合作发生了什么"),
    ("语义经验", "哪些规律可以复用"),
    ("关系模型", "用户重视什么，如何配合"),
    ("自我模型", "我擅长什么，怎样更像自己"),
]
for i, (title, desc) in enumerate(mems):
    y = Inches(3.2) + i * Inches(0.9)
    mem_bg = add_rounded_rect(slide, Inches(9.3), y, Inches(3.5), Inches(0.8),
                              fill_color=RGBColor(0x1a, 0x30, 0x2e),
                              border_color=RGBColor(0x30, 0x50, 0x48))
    add_text_box(slide, Inches(9.5), y + Inches(0.05), Inches(3.1), Inches(0.3),
                 title, size=11, color=RGBColor(0x8d, 0xd7, 0xca), bold=True)
    add_text_box(slide, Inches(9.5), y + Inches(0.35), Inches(3.1), Inches(0.35),
                 desc, size=9, color=RGBColor(0xc6, 0xd1, 0xdc))

# 对比框 - 底部
# 普通总结
cmp_bg1 = add_rounded_rect(slide, Inches(0.6), Inches(6.0), Inches(5.8), Inches(1.0),
                           fill_color=RGBColor(0x2a, 0x33, 0x3e),
                           border_color=RGBColor(0x44, 0x4f, 0x5c))
add_text_box(slide, Inches(0.8), Inches(6.05), Inches(5.4), Inches(0.25),
             "普通经验总结", size=9, color=DARK_MUTED, bold=True)
add_text_box(slide, Inches(0.8), Inches(6.3), Inches(5.4), Inches(0.5),
             "\"用户要求 PRD 包含背景和具体设计，以后沿用这个格式。\"",
             size=10, color=DARK_MUTED)

# 做梦机制
cmp_bg2 = add_rounded_rect(slide, Inches(6.8), Inches(6.0), Inches(5.8), Inches(1.0),
                           fill_color=RGBColor(0x1a, 0x30, 0x2e),
                           border_color=RGBColor(0x30, 0x50, 0x48))
add_text_box(slide, Inches(7.0), Inches(6.05), Inches(5.4), Inches(0.25),
             "做梦机制", size=9, color=RGBColor(0x8d, 0xd7, 0xca), bold=True)
add_text_box(slide, Inches(7.0), Inches(6.3), Inches(5.4), Inches(0.5),
             "\"多次合作让我发现：我的有效角色是帮助用户把想法说清楚。\"",
             size=10, color=RGBColor(0xc6, 0xd1, 0xdc))


# ══════════════════════════════════════════════════════════════
# SLIDE 6 — 闭环 + 目标
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, Inches(0), Inches(0), W, H, fill_color=NEAR_WHITE)

add_text_box(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(0.4),
             "05", size=14, color=BLUE, bold=True)
add_text_box(slide, Inches(0.6), Inches(0.9), Inches(12), Inches(0.6),
             "持续闭环：在线行动 ⇄ 离线做梦", size=26, color=INK, bold=True)
add_rect(slide, Inches(0.6), Inches(1.55), Inches(1.5), Pt(3), fill_color=CORAL)

# 四个步骤的水平流程 - 用卡片方式
steps_data = [
    ("① 在线行动", "双脑共同完成当前任务"),
    ("② 产生经历", "留下选择、反馈与结果"),
    ("③ 离线做梦", "重组经历并形成主体理解"),
    ("④ 带着过去再行动", "记忆进入下一次二阶观察"),
]

step_w = Inches(2.7)
step_h = Inches(1.6)
step_gap = Inches(0.35)
step_start_x = Inches(0.6)
step_start_y = Inches(2.3)

for i, (title, desc) in enumerate(steps_data):
    x = step_start_x + i * (step_w + step_gap)
    y = step_start_y
    
    # 卡片
    if i == 2:
        # 做梦用深色
        card = add_rounded_rect(slide, x, y, step_w, step_h,
                               fill_color=DARK_BG, border_color=RGBColor(0x33, 0x3f, 0x50))
        add_text_box(slide, x + Inches(0.15), y + Inches(0.15), step_w - Inches(0.3), Inches(0.4),
                     title, size=13, color=DARK_TEXT, bold=True)
        add_text_box(slide, x + Inches(0.15), y + Inches(0.65), step_w - Inches(0.3), Inches(0.7),
                     desc, size=10, color=DARK_MUTED)
    else:
        card = add_rounded_rect(slide, x, y, step_w, step_h,
                               fill_color=WHITE, border_color=RGBColor(0xdd, 0xdd, 0xdd))
        add_text_box(slide, x + Inches(0.15), y + Inches(0.15), step_w - Inches(0.3), Inches(0.4),
                     title, size=13, color=INK, bold=True)
        add_text_box(slide, x + Inches(0.15), y + Inches(0.65), step_w - Inches(0.3), Inches(0.7),
                     desc, size=10, color=MUTED)
    
    # 箭头 (except last)
    if i < 3:
        add_text_box(slide, x + step_w + Inches(0.02), y + Inches(0.4),
                     Inches(0.3), Inches(0.4), "→", size=20, color=BLUE, bold=True,
                     alignment=PP_ALIGN.CENTER)

# 闭环箭头（底部回环箭头）
add_text_box(slide, Inches(0.6), Inches(4.3), Inches(12), Inches(0.4),
             "← — — — — — — — — — — — — — — — — — — — — 闭环 ↺ — — — — — — — — — — — — — — — — — — — — →",
             size=11, color=BLUE, bold=True, alignment=PP_ALIGN.CENTER)

# 最终目标 - 大标题
add_rect(slide, Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.8),
         fill_color=BLUE_LIGHT, border_color=BLUE)

add_text_box(slide, Inches(1.0), Inches(5.4), Inches(11.3), Inches(0.5),
             "理论是起点，双脑是在线机制，做梦是离线机制",
             size=16, color=INK)

add_text_box(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.5),
             "主 体 连 续 性 —— 最终目标",
             size=24, color=BLUE, bold=True)

add_text_box(slide, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.4),
             "从一次性任务执行器，走向长期协作主体",
             size=13, color=MUTED)


# ── Save ──
output_path = "意识启发的Agent架构.pptx"
prs.save(output_path)
print(f"✅ 已生成: {output_path}")
